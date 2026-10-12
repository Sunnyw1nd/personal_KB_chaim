from __future__ import annotations

import json
from pathlib import Path

import pytest
from conftest import image_bytes
from openpyxl import load_workbook

from wechat_html_normalizer import normalize_wechat_article
from wechat_html_normalizer.image_annotation import generate, inventory, match_sources
from wechat_html_normalizer.media import sha256_file


def test_files_duplicates_failures_previews_and_human_data_survive_rerun(tmp_path: Path) -> None:
    pictures = tmp_path / "pictures"
    pictures.mkdir()
    nested = pictures / "子目录"
    nested.mkdir()
    (pictures / "=1+1.png").write_bytes(image_bytes("PNG", (200, 40)))
    (nested / "same.png").write_bytes((pictures / "=1+1.png").read_bytes())
    (pictures / "portrait.jpg").write_bytes(image_bytes("JPEG", (30, 200)))
    (pictures / "no-extension").write_bytes(image_bytes("WEBP"))
    (pictures / "broken.webp").write_bytes(b"not an image")
    before = {
        p.relative_to(pictures).as_posix(): sha256_file(p)
        for p in pictures.rglob("*")
        if p.is_file()
    }
    run = generate(pictures, tmp_path / "output", [])
    report = json.loads((run / "inventory.json").read_text(encoding="utf-8"))
    assert report["statistics"]["files"] == 5
    assert report["statistics"]["uniqueImages"] == 4
    assert report["statistics"]["duplicateGroups"] == 1
    assert report["statistics"]["previews"] == 4
    assert report["statistics"]["failed"] == 1
    workbook = run / "图片辅助标注.xlsx"
    book = load_workbook(workbook)
    assert book.sheetnames == ["标注主表", "选项", "文件与来源", "来源关系", "填写说明"]
    assert book["选项"].sheet_state == "hidden"
    sheet = book["标注主表"]
    assert len(sheet._images) == 4
    assert len(sheet.data_validations.dataValidation) == 4
    assert sheet.freeze_panes == "D5"
    assert sheet.auto_filter.ref == "A4:L9"
    for record in report["files"]:
        row = record["ordinal"] + 4
        assert sheet.cell(row, 3).value == record["fileName"]
        assert sheet.cell(row, 3).data_type == "s"
        assert sheet.cell(row, 5).hyperlink.target == (pictures / record["relativePath"]).as_uri()
        assert sheet.cell(row, 6).value == "待标注"
        assert sheet.cell(row, 8).value == "待标注"
        assert sheet.cell(row, 10).value is None
        assert sheet.cell(row, 11).value == "未审核"
        if record["decodable"]:
            image = next(i for i in sheet._images if i.anchor._from.row == row - 1)
            assert image.anchor.editAs == "twoCell"
            assert image._data() == (run / "previews" / f"{record['imageId']}.png").read_bytes()
    book.close()
    # Re-open because openpyxl image byte inspection consumes its internal streams.
    book = load_workbook(workbook)
    book["标注主表"]["J5"] = "人工已填写事实，必须保留"
    book.save(workbook)
    book.close()
    old_bytes = workbook.read_bytes()
    second = generate(pictures, tmp_path / "output", [])
    assert second != run
    assert workbook.read_bytes() == old_bytes
    assert {
        p.relative_to(pictures).as_posix(): sha256_file(p)
        for p in pictures.rglob("*")
        if p.is_file()
    } == before


def test_source_matches_are_by_hash_and_corruption_cannot_fall_back_to_manifest(
    saved_article: tuple[Path, Path, Path],
    tmp_path: Path,
) -> None:
    html, resources, output = saved_article
    result = normalize_wechat_article(html, resources, output)
    original_hash = sha256_file(resources / "640")
    matches, report = match_sources([output], {original_hash})
    assert report["documents"] == 1
    assert len(matches) == 1
    assert matches[0]["method"] == "verified_binary_sha256"
    assert matches[0]["assetId"] == result.document.assets[0].asset_id
    asset_path = result.output_directory / result.document.assets[0].output_path
    asset_path.write_bytes(image_bytes("PNG", (12, 12)))
    matches, report = match_sources([output], {original_hash})
    assert matches == []
    assert any(issue["code"] == "source_hash_mismatch" for issue in report["issues"])


def test_unsafe_overlap_and_symlinks_do_not_read_outside_inputs(tmp_path: Path) -> None:
    pictures = tmp_path / "pictures"
    pictures.mkdir()
    with pytest.raises(ValueError, match="overlap"):
        generate(pictures, pictures / "output", [])
    secret = tmp_path / "outside.png"
    secret.write_bytes(image_bytes("PNG"))
    try:
        (pictures / "linked.png").symlink_to(secret)
    except OSError:
        pytest.skip("Environment does not permit creating test symlinks.")
    with pytest.raises(ValueError, match="Linked input"):
        inventory(pictures)
