"""Offline image inventory and human annotation workbook; independent of import contracts."""

from __future__ import annotations

import argparse
import json
import os
import re
import warnings
from collections import Counter
from io import BytesIO
from pathlib import Path

from PIL import Image, ImageOps

from .media import sha256_file
from .models import Document, ResourceManifest

TYPES = [
    "待标注",
    "结构／流程图",
    "表格",
    "坐标／统计图",
    "代码截图",
    "界面截图",
    "普通文字图",
    "装饰文字",
    "二维码",
    "照片／场景图",
    "Logo／标志",
    "其他",
]
VALUES = ["待标注", "无", "低", "中", "高", "待上下文判断"]
METHODS = [
    "待判断",
    "仅归档／不进入检索",
    "OCR",
    "视觉描述",
    "OCR＋视觉理解",
    "表格提取",
    "图表理解",
    "代码转写",
    "二维码解码",
    "人工复核",
]
STATES = ["未审核", "待补充", "已确认"]
DEFAULT_SOURCE_ROOTS = [
    "output/documents",
    "output/real-samples",
    "output/managed-samples",
    "output/real-completion-pair",
]
MAX_FILES = 2000
MAX_FILE_BYTES = 50 * 1024 * 1024
PREVIEW_SIZE = (280, 170)


def _linked(path: Path) -> bool:
    return path.is_symlink() or (hasattr(path, "is_junction") and path.is_junction())


def _safe_path(path: Path) -> Path:
    path = path.absolute()
    if any(_linked(part) for part in [path, *path.parents]):
        raise ValueError("Symbolic links and junctions are not accepted: " + path.name)
    return path.resolve()


def _files(root: Path, notices: list[dict]) -> list[Path]:
    result = []
    for directory, dirs, names in os.walk(root, followlinks=False):
        for name in list(dirs):
            entry = Path(directory) / name
            if _linked(entry):
                dirs.remove(name)
                notices.append({"path": entry.relative_to(root).as_posix(), "code": "link_skipped"})
        for name in names:
            entry = Path(directory) / name
            if _linked(entry):
                raise ValueError(
                    "Linked input file rejected: " + entry.relative_to(root).as_posix()
                )
            if not entry.is_file():
                raise ValueError("Non-regular input file: " + name)
            result.append(entry)
            if len(result) > MAX_FILES:
                raise ValueError("File count exceeds limit; narrow the input/source directory.")
    return sorted(result, key=lambda p: (p.relative_to(root).as_posix().casefold(), str(p)))


def _preview(path: Path) -> tuple[dict, bytes | None]:
    info = {"format": None, "width": None, "height": None, "decodable": False, "error": None}
    if path.stat().st_size > MAX_FILE_BYTES:
        return {**info, "error": "file_size_limit"}, None
    try:
        with warnings.catch_warnings():
            warnings.simplefilter("error", Image.DecompressionBombWarning)
            with Image.open(path) as picture:
                info.update(format=picture.format, width=picture.width, height=picture.height)
                if picture.width * picture.height > 40_000_000:
                    raise ValueError("pixel_limit")
                picture.verify()
            with Image.open(path) as picture:
                picture.load()
                thumb = ImageOps.exif_transpose(picture).convert("RGBA")
                thumb.thumbnail(PREVIEW_SIZE, Image.Resampling.LANCZOS)
                background = Image.new("RGBA", PREVIEW_SIZE, (245, 247, 250, 255))
                background.alpha_composite(
                    thumb,
                    ((PREVIEW_SIZE[0] - thumb.width) // 2, (PREVIEW_SIZE[1] - thumb.height) // 2),
                )
                stream = BytesIO()
                background.convert("RGB").save(stream, format="PNG")
                info["decodable"] = True
                return info, stream.getvalue()
    except (
        OSError,
        ValueError,
        Image.DecompressionBombError,
        Image.DecompressionBombWarning,
    ) as exc:
        info["error"] = type(exc).__name__ + ": image decode failed or exceeds limits"
        return info, None


def inventory(root: Path) -> tuple[list[dict], dict[str, bytes], list[dict]]:
    notices: list[dict] = []
    rows, previews = [], {}
    for index, path in enumerate(_files(root, notices), 1):
        digest = sha256_file(path)
        info, preview = _preview(path)
        row = {
            "ordinal": index,
            "imageId": digest,
            "shortId": digest[:12],
            "relativePath": path.relative_to(root).as_posix(),
            "fileName": path.name,
            "sizeBytes": path.stat().st_size,
            **info,
        }
        suffix_format = {".jpg": "JPEG", ".jpeg": "JPEG", ".png": "PNG", ".webp": "WEBP"}.get(
            path.suffix.lower()
        )
        row["warning"] = (
            "extension_format_mismatch"
            if suffix_format and info["format"] and suffix_format != info["format"]
            else None
        )
        if preview:
            previews[digest] = preview
        rows.append(row)
    counts = Counter(row["imageId"] for row in rows)
    for row in rows:
        row["duplicateCount"] = counts[row["imageId"]]
        row["duplicateGroup"] = row["imageId"] if row["duplicateCount"] > 1 else None
    return rows, previews, notices


def match_sources(roots: list[Path], wanted: set[str]) -> tuple[list[dict], dict]:
    """Match recorded binary hashes, verifying exported assets where locally available."""
    relations, issues, seen = [], [], set()
    documents = 0
    existing_roots = 0
    for root_arg in roots:
        root = _safe_path(root_arg)
        if not root.is_dir():
            issues.append({"root": str(root), "code": "source_root_missing"})
            continue
        existing_roots += 1
        files = _files(root, issues)
        for path in files:
            if path.name != "document.json" or path in seen:
                continue
            seen.add(path)
            try:
                if path.stat().st_size > 10 * 1024 * 1024:
                    raise ValueError("source_json_size_limit")
                doc = Document.model_validate_json(path.read_text(encoding="utf-8"))
            except (ValueError, OSError):
                issues.append({"path": str(path), "code": "invalid_document"})
                continue
            documents += 1
            base = {
                "documentId": doc.document_id,
                "revisionId": doc.revision.revision_id if doc.revision else None,
                "title": doc.metadata.title,
                "documentJson": str(path),
            }
            body_hashes = {asset.sha256 for asset in doc.assets if asset.sha256}
            for asset in doc.assets:
                digest, method, availability = (
                    asset.sha256,
                    "recorded_binary_sha256",
                    "source_missing",
                )
                if asset.output_path:
                    candidate = path.parent / asset.output_path
                    try:
                        safe = _safe_path(candidate)
                        if not safe.is_relative_to(path.parent.resolve()):
                            raise ValueError("asset_outside_document")
                        if safe.is_file():
                            if safe.stat().st_size > MAX_FILE_BYTES:
                                raise ValueError("source_asset_size_limit")
                            actual = sha256_file(safe)
                            if digest and actual != digest:
                                issues.append(
                                    {
                                        "path": str(path),
                                        "assetId": asset.asset_id,
                                        "code": "source_hash_mismatch",
                                    }
                                )
                                # Never attach the recorded identity to a changed binary.
                                continue
                            digest, method, availability = (
                                actual,
                                "verified_binary_sha256",
                                "available",
                            )
                    except (OSError, ValueError):
                        issues.append(
                            {
                                "path": str(path),
                                "assetId": asset.asset_id,
                                "code": "unsafe_or_unreadable_source_asset",
                            }
                        )
                        continue
                if digest:
                    body_hashes.add(digest)
                if digest in wanted:
                    relations.append(
                        {
                            **base,
                            "imageId": digest,
                            "assetId": asset.asset_id,
                            "resourcePath": asset.output_path or asset.source_path,
                            "role": asset.role,
                            "method": method,
                            "availability": availability,
                        }
                    )
            manifest_path = path.parent / "resource-manifest.json"
            if manifest_path.exists():
                try:
                    _safe_path(manifest_path)
                    if manifest_path.stat().st_size > 10 * 1024 * 1024:
                        raise ValueError("manifest_size_limit")
                    manifest = ResourceManifest.model_validate_json(
                        manifest_path.read_text(encoding="utf-8")
                    )
                    if manifest.document_id != doc.document_id:
                        raise ValueError("manifest_document_mismatch")
                    for resource in manifest.resources:
                        if resource.sha256 in wanted and resource.sha256 not in body_hashes:
                            relations.append(
                                {
                                    **base,
                                    "imageId": resource.sha256,
                                    "assetId": None,
                                    "resourcePath": resource.path,
                                    "role": resource.reason_code,
                                    "method": "manifest_recorded_binary_sha256",
                                    "availability": "original_resource_not_checked",
                                }
                            )
                except (OSError, ValueError):
                    issues.append({"path": str(manifest_path), "code": "invalid_manifest"})
    return relations, {
        "roots": [str(p.resolve()) for p in roots],
        "documents": documents,
        "existingRoots": existing_roots,
        "issues": issues,
        "status": "not_executed" if not roots else "completed",
        "fallbackScope": "exported body assets only; raw caches not scanned",
    }


def _text(cell, value) -> None:
    if isinstance(value, str):
        # Explicit string cells prevent formula injection; clean XML control characters only.
        value = re.sub(r"[\x00-\x08\x0b\x0c\x0e-\x1f]", "�", value)
        if len(value) > 32767:
            value = value[:32740] + "…（完整值见清单）"
        cell.value = value
        cell.data_type = "s"
    else:
        cell.value = value


def _sheet(book, name: str, title: str, note: str, headers: list[str], widths: list[int]):
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter

    sheet = book.create_sheet(name)
    sheet.sheet_view.showGridLines = False
    sheet.sheet_view.zoomScale = 85
    for row, text in [(1, title), (2, note)]:
        sheet.merge_cells(start_row=row, start_column=1, end_row=row, end_column=len(headers))
        _text(sheet.cell(row, 1), text)
        sheet.cell(row, 1).font = Font(
            name="Microsoft YaHei", size=16 if row == 1 else 10, bold=row == 1, color="17365D"
        )
        sheet.cell(row, 1).alignment = Alignment(wrap_text=True, vertical="center")
        sheet.row_dimensions[row].height = 30
    for index, (header, width) in enumerate(zip(headers, widths), 1):
        cell = sheet.cell(4, index)
        _text(cell, header)
        cell.font = Font(name="Microsoft YaHei", size=10, bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="244765")
        cell.alignment = Alignment(wrap_text=True, vertical="center")
        sheet.column_dimensions[get_column_letter(index)].width = width
    sheet.row_dimensions[4].height = 32
    sheet.freeze_panes = "D5"
    return sheet


def _row(sheet, values: list, row: int, height: int = 36):
    from openpyxl.styles import Alignment, Font, PatternFill

    for col, value in enumerate(values, 1):
        cell = sheet.cell(row, col)
        _text(cell, value)
        cell.font = Font(name="Microsoft YaHei", size=10, color="273746")
        cell.alignment = Alignment(wrap_text=True, vertical="center")
        cell.fill = PatternFill("solid", fgColor="F3F6F9" if row % 2 else "FFFFFF")
    sheet.row_dimensions[row].height = height


def write_workbook(
    target: Path,
    root: Path,
    rows: list[dict],
    previews: dict[str, bytes],
    relations: list[dict],
    report: dict,
) -> None:
    from openpyxl import Workbook
    from openpyxl.drawing.image import Image as ExcelImage
    from openpyxl.drawing.spreadsheet_drawing import AnchorMarker, TwoCellAnchor
    from openpyxl.styles import Font, PatternFill
    from openpyxl.worksheet.datavalidation import DataValidation

    book = Workbook()
    book.remove(book.active)
    sheet = _sheet(
        book,
        "标注主表",
        "图片辅助标注 · 第一遍分类，第二遍补充事实",
        f"{len(rows)} 个文件；蓝色区域供人工填写。请勿排序或移动行：浮动缩略图尚未在 Excel/WPS 验证。",
        [
            "序号",
            "短 imageId",
            "原文件名",
            "缩略图",
            "打开原图",
            "图片主类型",
            "辅助标签",
            "知识价值",
            "推荐处理方式",
            "人工参考答案／关键事实",
            "审核状态",
            "备注",
        ],
        [6, 15, 26, 41, 12, 19, 20, 16, 23, 42, 13, 26],
    )
    for row in rows:
        index = row["ordinal"] + 4
        _row(
            sheet,
            [
                row["ordinal"],
                row["shortId"],
                row["fileName"],
                None if row["decodable"] else "预览失败，详见文件与来源",
                "打开原图",
                "待标注",
                None,
                "待标注",
                "待判断",
                None,
                "未审核",
                row["error"] or row["warning"],
            ],
            index,
            138,
        )
        for col in range(6, 13):
            sheet.cell(index, col).fill = PatternFill("solid", fgColor="EAF3FF")
        sheet.cell(index, 5).hyperlink = (root / row["relativePath"]).as_uri()
        sheet.cell(index, 5).font = Font(color="0563C1", underline="single", size=10)
        if row["decodable"]:
            image = ExcelImage(BytesIO(previews[row["imageId"]]))
            # Within-row two-cell anchors; intended to move/size with cells, not a sort guarantee.
            image.anchor = TwoCellAnchor(
                editAs="twoCell",
                _from=AnchorMarker(col=3, row=index - 1, colOff=38100, rowOff=38100),
                to=AnchorMarker(col=3, row=index - 1, colOff=2705100, rowOff=1657350),
            )
            sheet.add_image(image)
    options = book.create_sheet("选项")
    for col, choices in enumerate([TYPES, VALUES, METHODS, STATES], 1):
        for row, choice in enumerate(choices, 1):
            _text(options.cell(row, col), choice)
        if rows:
            letter = "ABCD"[col - 1]
            validation = DataValidation(
                type="list", formula1=f"'选项'!${letter}$1:${letter}${len(choices)}"
            )
            validation.allow_blank = True
            validation.showErrorMessage = True
            validation.errorTitle = "请选择下拉选项"
            validation.error = "混合需求请写入辅助标签或备注。"
            sheet.add_data_validation(validation)
            validation.add(
                f"{['F', 'H', 'I', 'K'][col - 1]}5:{['F', 'H', 'I', 'K'][col - 1]}{len(rows) + 4}"
            )
    options.sheet_state = "hidden"
    files = _sheet(
        book,
        "文件与来源",
        "文件清单 · 一文件一行，完整身份可追溯",
        "来源关系按完整二进制 SHA-256 精确匹配；一张图片的多条关系见“来源关系”。",
        [
            "序号",
            "完整 imageId / SHA-256",
            "相对路径",
            "原文件名",
            "实际格式",
            "宽",
            "高",
            "字节数",
            "解码状态",
            "异常",
            "重复组（完整哈希）",
            "组文件数",
            "来源匹配状态",
            "关系数",
        ],
        [7, 68, 45, 42, 12, 10, 10, 14, 12, 26, 68, 12, 20, 10],
    )
    for row in rows:
        _row(
            files,
            [
                row["ordinal"],
                row["imageId"],
                row["relativePath"],
                row["fileName"],
                row["format"],
                row["width"],
                row["height"],
                row["sizeBytes"],
                "成功" if row["decodable"] else "失败",
                row["error"] or row["warning"],
                row["duplicateGroup"],
                row["duplicateCount"],
                row["sourceStatus"],
                row["relationCount"],
            ],
            row["ordinal"] + 4,
        )
    links = _sheet(
        book,
        "来源关系",
        "精确来源关系 · 一关系一行",
        "同一关系在不同输出副本中出现时保留证据路径。旧输出没有 revisionId 时留空。",
        [
            "完整 imageId",
            "documentId",
            "revisionId",
            "assetId",
            "标题",
            "资源路径",
            "资源角色",
            "匹配依据",
            "资源可访问状态",
            "来源 JSON",
        ],
        [68, 30, 30, 24, 45, 45, 26, 34, 30, 70],
    )
    keys = [
        "imageId",
        "documentId",
        "revisionId",
        "assetId",
        "title",
        "resourcePath",
        "role",
        "method",
        "availability",
        "documentJson",
    ]
    for index, relation in enumerate(relations, 5):
        _row(links, [relation.get(key) for key in keys], index, 45)
    guide = _sheet(
        book,
        "填写说明",
        "填写说明 · 人工确认后才可成为参考事实",
        "当前交付为辅助标注工具；分类和参考事实尚未审核，不代表评测基准或识别方案已完成。",
        ["项目", "说明"],
        [28, 115],
    )
    instructions = [
        (
            "第一遍",
            "完成所有可用图片的主类型、知识价值和推荐处理方式；辅助标签用分号分隔。蓝色为填写区。",
        ),
        (
            "第二遍",
            "仅对需要识别的图片填写关键参考事实。含糊、需要正文上下文或看不清时说明不确定性，并选待补充。",
        ),
        (
            "审核标准",
            "逐图核对分类、价值、处理方式及适用的参考事实后才选已确认；未审核与无知识价值不是一回事。",
        ),
        ("主类型选项", "；".join(TYPES)),
        ("知识价值选项", "；".join(VALUES)),
        ("处理方式选项", "；".join(METHODS)),
        ("审核状态", "；".join(STATES)),
        (
            "辅助标签示例",
            "复杂布局；含文字；含卡通人物；含方向关系；混合内容；低分辨率。自由文本，不使用宏多选。",
        ),
        (
            "通用建议（非机器判图）",
            "结构图→视觉描述或OCR＋视觉理解；表格→表格提取；坐标图→图表理解；代码→代码转写；文字→OCR；混合截图→人工复核。",
        ),
        (
            "上下文与价值",
            "照片、Logo、二维码和装饰文字不能仅凭外观断定无价值；缺上下文时选待上下文判断。",
        ),
        ("示例：装饰图", "仅作装饰，没有需要提取的知识内容。（必须经用户确认，不能自动套用。）"),
        ("示例：结构图", "记录关键节点、分组、箭头方向与关系；可列出不允许推断的关系。"),
        ("示例：表格", "表头、重点单元格、行列及合并关系。"),
        ("示例：坐标图", "轴名、单位、图例、趋势，以及能清楚读取的关键数值。"),
        ("示例：代码", "若只测路由，可暂不录全文；完整转写评估需要另补完整参考代码。"),
        ("示例：二维码", "未来使用解码工具取得实际内容后人工确认；不要视觉猜测地址，无需手抄。"),
        ("示例：文字／界面", "记录关键文字、区域或状态；看不清部分保留不确定性。"),
        (
            "评测边界",
            "关键事实只用于事实覆盖和错误检查，不能直接宣称完整OCR字符准确率。机器草稿必须对照原图审核；被测模型输出不能直接成为标准答案。",
        ),
        (
            "原图链接",
            "本地file URI指向生成时的绝对路径；Excel/WPS可能要求确认打开。移动原图后链接可能失效，内嵌缩略图仍保留。",
        ),
        (
            "筛选／排序限制",
            "表头提供筛选；图片使用随单元格移动的浮动锚点，但未在桌面Excel/WPS验证。请勿排序、插入或移动行；筛选后如图片位置异常，恢复全表并以序号、文件名和原图链接核对。",
        ),
        (
            "缩略图",
            "等比例缩放、EXIF方向校正、透明底铺浅灰色；长图细节请打开原图。动画仅预览首帧。",
        ),
        (
            "身份与重复",
            "imageId为文件完整SHA-256；短ID仅供展示。同二进制不同文件各占一行，重复组不删除。将来划分数据集时同组应归入同一集合。",
        ),
        (
            "来源匹配",
            "verified_binary_sha256=复算输出资产；recorded_binary_sha256=记录哈希相同但输出资产缺失；manifest_recorded_binary_sha256=清单哈希，原缓存未复算。空assetId表示无正式资产ID，不能根据文件名猜测。",
        ),
        (
            "扫描范围",
            "输入递归扫描所有普通文件，不按扩展名忽略损坏文件；不跟随符号链接/目录联接；输出禁止放在输入目录中。",
        ),
        ("来源扫描目录", "\n".join(report["sources"]["roots"]) or "未执行来源匹配"),
        ("来源扫描问题", json.dumps(report["sources"]["issues"], ensure_ascii=False)),
        (
            "再次生成",
            "默认创建新的编号目录，已有标注文件不会覆盖。新工作簿不会自动合并旧的人工填写。完整清单为同目录inventory.json。",
        ),
    ]
    for index, (label, content) in enumerate(instructions, 5):
        _row(guide, [label, content], index, 60 if len(content) > 100 else 42)
    guide.freeze_panes = "B5"
    for tab in [sheet, files, links]:
        if tab.max_row >= 5:
            tab.auto_filter.ref = f"A4:{tab.cell(tab.max_row, tab.max_column).coordinate}"
    book.active = 0
    book.save(target)


def generate(input_root: Path, output_root: Path, source_roots: list[Path]) -> Path:
    input_root, output_root = _safe_path(input_root), _safe_path(output_root)
    if not input_root.is_dir():
        raise ValueError("Input directory does not exist.")
    if output_root.is_relative_to(input_root) or input_root.is_relative_to(output_root):
        raise ValueError("Input and output directories must not overlap.")
    rows, previews, notices = inventory(input_root)
    relations, sources = match_sources(source_roots, {row["imageId"] for row in rows})
    counts = Counter(relation["imageId"] for relation in relations)
    for row in rows:
        count = counts[row["imageId"]]
        row["relationCount"] = count
        row["sourceStatus"] = (
            "多个匹配"
            if count > 1
            else "精确匹配"
            if count
            else "未执行"
            if not source_roots
            else "来源缺失"
            if not sources["documents"]
            else "未匹配"
        )
    statistics = {
        "files": len(rows),
        "uniqueImages": len({r["imageId"] for r in rows}),
        "duplicateGroups": len({r["duplicateGroup"] for r in rows if r["duplicateGroup"]}),
        "previews": sum(r["decodable"] for r in rows),
        "failed": sum(not r["decodable"] for r in rows),
        "extensionMismatches": sum(bool(r["warning"]) for r in rows),
        "formats": dict(Counter(r["format"] or "unknown" for r in rows)),
        "matchedFiles": sum(r["relationCount"] > 0 for r in rows),
        "multipleMatchFiles": sum(r["relationCount"] > 1 for r in rows),
        "relations": len(relations),
    }
    output_root.mkdir(parents=True, exist_ok=True)
    number = 1
    while True:
        run = output_root / f"annotation-{number:03d}"
        try:
            run.mkdir()
            break
        except FileExistsError:
            number += 1
    report = {
        "schemaVersion": "1.0.0",
        "inputRoot": str(input_root),
        "statistics": statistics,
        "sources": sources,
        "scanNotices": notices,
        "files": rows,
        "relations": relations,
        "status": "complete",
        "originalInputsUnchanged": True,
    }
    preview_dir = run / "previews"
    preview_dir.mkdir()
    for digest, content in previews.items():
        (preview_dir / f"{digest}.png").write_bytes(content)
    temporary = run / "workbook.pending.xlsx"
    write_workbook(temporary, input_root, rows, previews, relations, report)
    before = {row["relativePath"]: row["imageId"] for row in rows}
    after = {p.relative_to(input_root).as_posix(): sha256_file(p) for p in _files(input_root, [])}
    if before != after:
        raise RuntimeError("Input files changed during generation; workbook not published.")
    from openpyxl import load_workbook

    check = load_workbook(temporary)
    assert check["标注主表"].max_row == len(rows) + 4
    assert len(check["标注主表"]._images) == statistics["previews"]
    check.close()
    temporary.rename(run / "图片辅助标注.xlsx")
    (run / "inventory.json").write_text(
        json.dumps(report, ensure_ascii=False, indent=2, sort_keys=True) + "\n", encoding="utf-8"
    )
    return run


def main() -> None:
    parser = argparse.ArgumentParser(description="Generate a new offline image annotation XLSX.")
    parser.add_argument("--input", type=Path, default=Path("test_pictures"))
    parser.add_argument("--output", type=Path, default=Path("output/image-annotation"))
    parser.add_argument(
        "--source-root",
        action="append",
        type=Path,
        help="Repeat for multiple bounded document output roots.",
    )
    parser.add_argument("--no-source-match", action="store_true")
    args = parser.parse_args()
    roots = (
        [] if args.no_source_match else args.source_root or list(map(Path, DEFAULT_SOURCE_ROOTS))
    )
    try:
        result = generate(args.input, args.output, roots)
    except (OSError, ValueError, RuntimeError) as exc:
        parser.exit(2, f"Annotation generation failed: {exc}\n")
    print(json.dumps({"outputDirectory": str(result)}, ensure_ascii=True))


if __name__ == "__main__":
    main()
