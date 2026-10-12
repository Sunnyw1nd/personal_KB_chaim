from __future__ import annotations

import json
from collections import Counter, defaultdict
from datetime import UTC, datetime
from pathlib import Path

from ..media import sha256_file
from .models import BatchRunManifest, EnrichmentRunRecord, RunStatus
from .security import redact_sensitive_text

TYPE_MAP = {
    "结构／流程图": "architecture_flow",
    "表格": "table",
    "坐标／统计图": "chart",
    "代码截图": "code",
    "界面截图": "ui_screenshot",
    "普通文字图": "text_image",
    "装饰文字": "decorative_text",
    "二维码": "qr_code",
    "照片／场景图": "photo_scene",
    "Logo／标志": "logo_mark",
    "其他": "other",
}
VALUE_MAP = {
    "无": "none",
    "低": "low",
    "中": "medium",
    "高": "high",
    "待上下文判断": "context_dependent",
}
ACTION_MAP = {
    "仅归档／不进入检索": "exclude_from_retrieval",
    "OCR": "transcribe_text",
    "视觉描述": "describe",
    "OCR＋视觉理解": "ocr_and_vision",
    "表格提取": "extract_table",
    "图表理解": "understand_chart",
    "代码转写": "transcribe_code",
    "二维码解码": "decode_qr",
    "人工复核": "manual_review",
}


def _header(sheet) -> dict[str, int]:
    return {cell.value: cell.column for cell in sheet[4] if isinstance(cell.value, str)}


def read_benchmark(workbook_path: Path, image_root: Path) -> tuple[list[dict], dict]:
    try:
        from openpyxl import load_workbook
    except ImportError as exc:  # pragma: no cover
        raise RuntimeError('install the "annotation" or "enrichment" optional dependency') from exc
    workbook_path = workbook_path.resolve(strict=True)
    image_root = image_root.resolve(strict=True)
    workbook_sha256 = sha256_file(workbook_path)
    book = load_workbook(workbook_path, read_only=True, data_only=False)
    try:
        main = book["标注主表"]
        files = book["文件与来源"]
        main_header, files_header = _header(main), _header(files)
        required_main = {
            "序号",
            "图片主类型",
            "知识价值",
            "推荐处理方式",
            "人工参考答案／关键事实",
            "审核状态",
            "备注",
        }
        required_files = {"序号", "完整 imageId / SHA-256", "相对路径", "原文件名"}
        if not required_main <= main_header.keys() or not required_files <= files_header.keys():
            raise ValueError("benchmark workbook is missing required columns")
        metadata: dict[int, dict] = {}
        for row in range(5, files.max_row + 1):
            ordinal = files.cell(row, files_header["序号"]).value
            if not isinstance(ordinal, int):
                continue
            digest = files.cell(row, files_header["完整 imageId / SHA-256"]).value
            relative = files.cell(row, files_header["相对路径"]).value
            filename = files.cell(row, files_header["原文件名"]).value
            if not isinstance(digest, str) or not isinstance(relative, str):
                raise TypeError(f"invalid file identity at workbook row {row}")
            candidate = (image_root / Path(*relative.split("/"))).resolve(strict=True)
            if not candidate.is_relative_to(image_root) or candidate.name != filename:
                raise ValueError(f"unsafe or mismatched benchmark path at row {row}")
            if sha256_file(candidate) != digest:
                raise ValueError(f"benchmark image hash mismatch at row {row}")
            metadata[ordinal] = {
                "imageSha256": digest,
                "relativePath": relative,
                "fileName": filename,
            }
        rows = []
        for row in range(5, main.max_row + 1):
            ordinal = main.cell(row, main_header["序号"]).value
            if not isinstance(ordinal, int):
                continue
            if ordinal not in metadata:
                raise ValueError(f"main row {row} has no matching file metadata")
            image_type_raw = main.cell(row, main_header["图片主类型"]).value
            value_raw = main.cell(row, main_header["知识价值"]).value
            action_raw = main.cell(row, main_header["推荐处理方式"]).value
            review_status = main.cell(row, main_header["审核状态"]).value
            raw_reference = main.cell(row, main_header["人工参考答案／关键事实"]).value
            reference = raw_reference
            reference_redacted = False
            if isinstance(raw_reference, str):
                reference, reference_redacted = redact_sensitive_text(raw_reference)
            raw_note = main.cell(row, main_header["备注"]).value
            note = raw_note
            note_redacted = False
            if isinstance(raw_note, str):
                note, note_redacted = redact_sensitive_text(raw_note)
            rows.append(
                {
                    **metadata[ordinal],
                    "ordinal": ordinal,
                    "benchmark": {
                        "imageTypeRaw": image_type_raw,
                        "imageType": TYPE_MAP.get(image_type_raw),
                        "knowledgeValueRaw": value_raw,
                        "knowledgeValue": VALUE_MAP.get(value_raw),
                        "actionRaw": action_raw,
                        "action": ACTION_MAP.get(action_raw),
                        "retrievalDecision": _retrieval_from_action(action_raw),
                        "referenceAnswer": reference,
                        "sensitiveContentRedacted": reference_redacted or note_redacted,
                        "reviewStatus": review_status,
                        "confirmed": review_status == "已确认",
                        "note": note,
                    },
                }
            )
    finally:
        book.close()
    return rows, {
        "workbookFileName": workbook_path.name,
        "workbookSha256": workbook_sha256,
        "rows": len(rows),
        "uniqueImages": len({row["imageSha256"] for row in rows}),
        "reviewStates": dict(Counter(row["benchmark"]["reviewStatus"] for row in rows)),
        "confirmedRows": sum(row["benchmark"]["confirmed"] for row in rows),
        "referenceAnswersPresent": sum(bool(row["benchmark"]["referenceAnswer"]) for row in rows),
        "rowsWithSensitiveContentRedacted": sum(
            row["benchmark"]["sensitiveContentRedacted"] for row in rows
        ),
    }


def _retrieval_from_action(value: object) -> str | None:
    if value == "仅归档／不进入检索":
        return "exclude"
    if value == "人工复核":
        return "manual_review"
    if value in ACTION_MAP:
        return "include"
    return None


def _load_results(manifest_path: Path) -> tuple[BatchRunManifest, dict[str, EnrichmentRunRecord]]:
    manifest_path = manifest_path.resolve(strict=True)
    manifest = BatchRunManifest.model_validate_json(manifest_path.read_text(encoding="utf-8"))
    output_root = manifest_path.parents[2]
    results: dict[str, EnrichmentRunRecord] = {}
    for item in manifest.items:
        if item.status not in (RunStatus.SUCCESS, RunStatus.CACHED) or not item.record_path:
            continue
        path = (output_root / Path(*item.record_path.split("/"))).resolve(strict=True)
        if not path.is_relative_to(output_root):
            raise ValueError("run manifest record path escapes enrichment output")
        record = EnrichmentRunRecord.model_validate_json(path.read_text(encoding="utf-8"))
        if record.image_sha256 != item.image_sha256 or record.cache_key != item.cache_key:
            raise ValueError("run item and record identity do not match")
        results[item.image_sha256] = record
    return manifest, results


def _agreement(rows: list[dict], *, confirmed_only: bool) -> dict:
    fields = ["imageType", "knowledgeValue", "action", "retrievalDecision"]
    result = {}
    for field in fields:
        eligible = [
            row
            for row in rows
            if row["benchmark"].get(field) is not None
            and row["machine"] is not None
            and (row["benchmark"]["confirmed"] or not confirmed_only)
        ]
        matches = sum(row["benchmark"][field] == row["machine"].get(field) for row in eligible)
        result[field] = {
            "eligible": len(eligible),
            "matches": matches,
            "agreementRate": matches / len(eligible) if eligible else None,
        }
    return result


def evaluate_run(
    manifest_path: Path,
    workbook_path: Path,
    image_root: Path,
    output_root: Path,
) -> tuple[Path, dict]:
    benchmark, benchmark_summary = read_benchmark(workbook_path, image_root)
    manifest, results = _load_results(manifest_path)
    combined = []
    by_type: dict[str, Counter] = defaultdict(Counter)
    for row in benchmark:
        record = results.get(row["imageSha256"])
        machine = None
        if record and record.result:
            machine = record.result.model_dump(mode="json", by_alias=True, exclude_none=True)
            benchmark_type = row["benchmark"]["imageType"] or "unresolved"
            by_type[benchmark_type]["results"] += 1
            by_type[benchmark_type][
                "typeMatches"
                if machine["imageType"] == row["benchmark"]["imageType"]
                else "typeDiffers"
            ] += 1
        combined.append(
            {
                **row,
                "machine": machine,
                "manualQualityGrade": None,
                "manualQualityNote": None,
            }
        )
    statuses = Counter(item.status.value for item in manifest.items)
    elapsed = [
        results[row["imageSha256"]].elapsed_ms for row in benchmark if row["imageSha256"] in results
    ]
    report = {
        "schemaVersion": "1.0.0",
        "createdAt": datetime.now(UTC).isoformat().replace("+00:00", "Z"),
        "run": {
            "runId": manifest.run_id,
            "provider": manifest.provider,
            "requestedModel": manifest.requested_model,
            "actualModels": manifest.actual_models,
            "promptVersion": manifest.prompt_version,
            "promptSha256": manifest.prompt_sha256,
            "resultSchemaVersion": manifest.result_schema_version,
            "status": manifest.status,
            "usage": manifest.usage.model_dump(mode="json", by_alias=True, exclude_none=True)
            if manifest.usage
            else None,
        },
        "benchmark": benchmark_summary,
        "results": {
            "statusCounts": dict(statuses),
            "resultRecords": len(results),
            "missingResults": len(benchmark) - len(results),
            "elapsedMs": {
                "total": sum(elapsed),
                "minimum": min(elapsed) if elapsed else None,
                "maximum": max(elapsed) if elapsed else None,
                "average": sum(elapsed) / len(elapsed) if elapsed else None,
            },
        },
        "agreement": {
            "confirmedOnly": _agreement(combined, confirmed_only=True),
            "allResolvableRowsProvisional": _agreement(combined, confirmed_only=False),
            "note": (
                "Unconfirmed rows are reported as provisional raw agreement and are not treated "
                "as adjudicated model errors."
            ),
        },
        "byBenchmarkType": {key: dict(value) for key, value in sorted(by_type.items())},
        "qualityReview": {
            "pendingRows": len(results),
            "grades": [
                "no_change",
                "minor_edit",
                "substantive_edit",
                "unusable",
                "benchmark_review_needed",
            ],
            "note": (
                "Open-ended descriptions and key facts require human review; no string similarity "
                "score is presented as quality."
            ),
        },
        "items": combined,
    }
    output_root = output_root.resolve()
    output_root.mkdir(parents=True, exist_ok=True)
    stamp = datetime.now(UTC).strftime("%Y%m%dT%H%M%SZ")
    directory = output_root / f"{manifest.run_id}-{stamp}"
    if directory.exists():
        raise FileExistsError(f"evaluation output already exists: {directory}")
    directory.mkdir()
    json_path = directory / "baseline-evaluation.json"
    json_path.write_text(
        json.dumps(report, ensure_ascii=False, indent=2, sort_keys=True) + "\n",
        encoding="utf-8",
    )
    (directory / "baseline-evaluation.md").write_text(_markdown(report), encoding="utf-8")
    return json_path, report


def _markdown(report: dict) -> str:
    run, benchmark, results = report["run"], report["benchmark"], report["results"]
    lines = [
        "# 图片 Enrichment 基线评测",
        "",
        f"- Run：`{run['runId']}`",
        f"- Provider / model：`{run['provider']}` / `{run['requestedModel']}`",
        f"- Prompt / Schema：`{run['promptVersion']}` / `{run['resultSchemaVersion']}`",
        f"- 基准行数：{benchmark['rows']}；已确认：{benchmark['confirmedRows']}",
        f"- 结果记录：{results['resultRecords']}；缺少结果：{results['missingResults']}",
        f"- 状态：{json.dumps(results['statusCounts'], ensure_ascii=False, sort_keys=True)}",
        "",
        "## 标签符合度",
        "",
        "已确认行是正式统计口径；未确认行只显示原始、暂定符合度。",
        "",
        "| 字段 | 已确认 eligible / matches | 全部可比较 eligible / matches |",
        "| --- | ---: | ---: |",
    ]
    confirmed = report["agreement"]["confirmedOnly"]
    provisional = report["agreement"]["allResolvableRowsProvisional"]
    for field in ("imageType", "knowledgeValue", "action", "retrievalDecision"):
        lines.append(
            f"| {field} | {confirmed[field]['eligible']} / {confirmed[field]['matches']} | "
            f"{provisional[field]['eligible']} / {provisional[field]['matches']} |"
        )
    lines.extend(
        [
            "",
            "## 尚待人工评审",
            "",
            (
                f"{report['qualityReview']['pendingRows']} 条机器结果需要按原图核对关键事实、遗漏、"
                "无依据新增和关系方向。开放式内容没有用字符串相似度冒充质量分数。"
            ),
            "",
            "逐图机器结果、人工基准与空白工作量等级见 `baseline-evaluation.json`。",
        ]
    )
    return "\n".join(lines) + "\n"
